import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
from datetime import datetime
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

FIELDS = [
    ("CAD Incident Number", ["CAD Incident Number"]),
    ("Incident Date / Time", ["Incident Date / Time"]),
    ("Shift", ["Shift"]),
    ("Station", ["Station"]),
    ("Units", ["Units"]),
    ("Address", ["Address", "Location"]),
]

def clean_display(value, field):
    if value is None:
        return ""
    if field == "Units" and isinstance(value, str):
        value = re.sub(r"<br\s*/?>", "\n", value, flags=re.I)
    return value

def read_source(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}

    mapping = {}
    for out, choices in FIELDS:
        found = next((headers[x] for x in choices if x in headers), None)
        if found is None:
            raise ValueError(
                f'Missing required source column for "{out}". '
                f'Expected one of: {", ".join(choices)}'
            )
        mapping[out] = found

    rows = []
    for r in range(2, ws.max_row + 1):
        values = [
            clean_display(ws.cell(r, mapping[out]).value, out)
            for out, _ in FIELDS
        ]
        if any(v not in ("", None) for v in values):
            rows.append(values)
    return rows

def build_report(source, report_type, destination):
    rows = read_source(source)

    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions Report"

    title = f"{report_type.upper()} EXCEPTIONS REPORT"

    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:C2")
    ws["A2"] = f"Generated: {datetime.now():%B %d, %Y at %I:%M %p}"
    ws.merge_cells("D2:F2")
    ws["D2"] = f"Total Exceptions: {len(rows)}"

    for cell in ("A2", "D2"):
        ws[cell].font = Font(italic=True, color="404040")
        ws[cell].alignment = Alignment(
            horizontal="left" if cell == "A2" else "right"
        )

    header_row = 4
    headers = [field[0] for field in FIELDS]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[header_row].height = 24

    thin = Side(style="thin", color="D9E2F3")
    for r_idx, row in enumerate(rows, header_row + 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

    if rows:
        ref = f"A{header_row}:F{header_row + len(rows)}"
        table = Table(displayName="ExceptionsTable", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    widths = [22, 23, 12, 13, 30, 42]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"
    ws.oddFooter.center.text = "OFD Exceptions Report"
    ws.oddFooter.right.text = "Page &P of &N"

    wb.save(destination)
    return len(rows)

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("OFD Exceptions Report Generator")
        self.geometry("720x430")
        self.minsize(680, 400)
        self.configure(bg="#eef2f6")

        self.source = tk.StringVar()
        self.report_type = tk.StringVar(value="Draft Reports")
        self.status = tk.StringVar(value="Select a report type and source file.")

        style = ttk.Style(self)
        if "vista" in style.theme_names():
            style.theme_use("vista")

        style.configure(
            "Title.TLabel",
            font=("Segoe UI", 20, "bold"),
            foreground="#17365D",
            background="#eef2f6",
        )
        style.configure(
            "Sub.TLabel",
            font=("Segoe UI", 10),
            foreground="#555555",
            background="#eef2f6",
        )
        style.configure("Card.TFrame", background="white")
        style.configure("Card.TLabel", background="white", font=("Segoe UI", 10))
        style.configure("Action.TButton", font=("Segoe UI", 11, "bold"), padding=10)

        outer = ttk.Frame(self, padding=24)
        outer.pack(fill="both", expand=True)

        ttk.Label(
            outer,
            text="OFD Exceptions Report Generator",
            style="Title.TLabel"
        ).pack(anchor="w")

        ttk.Label(
            outer,
            text="Create a polished Excel report from your incomplete-report export.",
            style="Sub.TLabel",
        ).pack(anchor="w", pady=(2, 18))

        card = ttk.Frame(outer, style="Card.TFrame", padding=22)
        card.pack(fill="both", expand=True)

        ttk.Label(
            card,
            text="1. Exception report type",
            style="Card.TLabel"
        ).grid(row=0, column=0, columnspan=3, sticky="w")

        combo = ttk.Combobox(
            card,
            textvariable=self.report_type,
            values=["Draft Reports", "In Review Reports", "2nd Review Reports"],
            state="readonly",
            width=28,
        )
        combo.grid(row=1, column=0, columnspan=3, sticky="w", pady=(6, 18))

        ttk.Label(
            card,
            text="2. Source export",
            style="Card.TLabel"
        ).grid(row=2, column=0, columnspan=3, sticky="w")

        entry = ttk.Entry(card, textvariable=self.source)
        entry.grid(
            row=3,
            column=0,
            columnspan=2,
            sticky="ew",
            pady=(6, 18),
            ipady=5,
        )

        ttk.Button(
            card,
            text="Browse…",
            command=self.browse
        ).grid(
            row=3,
            column=2,
            sticky="e",
            padx=(10, 0),
            pady=(6, 18),
        )

        ttk.Separator(card).grid(
            row=4,
            column=0,
            columnspan=3,
            sticky="ew",
            pady=(2, 18),
        )

        ttk.Button(
            card,
            text="Generate Excel Report",
            style="Action.TButton",
            command=self.generate,
        ).grid(row=5, column=0, columnspan=3, sticky="ew")

        ttk.Label(
            card,
            textvariable=self.status,
            style="Card.TLabel",
            wraplength=600,
        ).grid(row=6, column=0, columnspan=3, sticky="w", pady=(16, 0))

        card.columnconfigure(0, weight=1)
        card.columnconfigure(1, weight=1)

    def browse(self):
        path = filedialog.askopenfilename(
            title="Select export",
            filetypes=[
                ("Excel files", "*.xlsx"),
                ("All files", "*.*"),
            ],
        )
        if path:
            self.source.set(path)
            self.status.set(f"Ready: {Path(path).name}")

    def generate(self):
        source = self.source.get().strip()

        if not source or not Path(source).exists():
            messagebox.showerror(
                "Source file required",
                "Select the Excel export first.",
            )
            return

        suggested = (
            f"{self.report_type.get().replace(' ', '_')}_Exceptions_"
            f"{datetime.now():%Y-%m-%d}.xlsx"
        )

        destination = filedialog.asksaveasfilename(
            title="Save exceptions report",
            defaultextension=".xlsx",
            initialfile=suggested,
            filetypes=[("Excel workbook", "*.xlsx")],
        )

        if not destination:
            return

        try:
            count = build_report(
                source,
                self.report_type.get(),
                destination,
            )
            self.status.set(
                f"Created {Path(destination).name} with {count} exceptions."
            )
            messagebox.showinfo(
                "Report created",
                f"Excel report created successfully.\n\n"
                f"{count} exceptions\n{destination}",
            )
        except Exception as exc:
            messagebox.showerror("Could not create report", str(exc))

if __name__ == "__main__":
    App().mainloop()
